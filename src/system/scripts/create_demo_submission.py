from __future__ import annotations

import os
import sys
from io import BytesIO
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "conference_project.settings")

import django
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document

django.setup()

from submissions.services import SubmissionService
from submissions.workflow import WorkflowEngine


def build_demo_docx() -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("УДК 004.8")
    document.add_paragraph("ПРИМЕНЕНИЕ НЕЙРОСЕТЕЙ ДЛЯ АНАЛИЗА УЧЕБНЫХ ТЕКСТОВ")
    document.add_paragraph("Иванов И.И., Петров П.П.")
    document.add_paragraph("ТГТУ")
    document.add_paragraph("ivanov@example.com")
    document.add_paragraph("Научный руководитель: Сидоров С.С.")
    document.add_paragraph("Аннотация: Демонстрационный материал для проверки интеграции ядра и DOCX-модуля.")
    document.add_paragraph("Ключевые слова: нейросети, образование, анализ текста")
    document.add_heading("Введение", level=1)
    document.add_paragraph("Цель работы — продемонстрировать полный маршрут обработки DOCX.")
    document.add_heading("Методика", level=1)
    document.add_paragraph("Материал загружается, анализируется и приводится к шаблону конференции.")
    document.add_heading("Список литературы", level=1)
    document.add_paragraph("1. Иванов И.И. Нейросетевые методы. 2025.")
    document.save(buffer)
    return buffer.getvalue()


service = SubmissionService()
submission = service.create_submission(
    {
        "author_contact": {
            "full_name": "Иванов Иван Иванович",
            "email": "ivanov@example.com",
            "organization": "ТГТУ",
        },
        "metadata": {
            "title_ru": "Применение нейросетей для анализа учебных текстов",
            "section": "Искусственный интеллект в образовании",
            "keywords_ru": ["нейросети", "образование", "анализ текста"],
            "abstract_ru": "Демонстрационная заявка для проверки интеграции.",
            "authors": [],
        },
    },
    uploaded_file=SimpleUploadedFile(
        "demo.docx",
        build_demo_docx(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ),
)

submission_id = submission["submission_id"]
results = WorkflowEngine().run(submission_id)
print(f"Заявка: {submission_id}")
for result in results:
    print(f"- {result['stage_id']}: {result['status']} — {result['message']}")
print(f"Итоговый статус: {service.get_submission(submission_id)['status']}")
