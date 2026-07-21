from __future__ import annotations

import json
import tempfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from docx import Document
from PIL import Image

from submissions.models import Organization, Submission, SubmissionFile, WorkflowStageResult
from submissions.plugin_registry import registry
from submissions.services import SubmissionService, resolve_stored_file_path
from submissions.workflow import WorkflowEngine


def build_valid_docx() -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("УДК 004.8")
    document.add_paragraph("ПРИМЕНЕНИЕ НЕЙРОСЕТЕЙ В ОБРАЗОВАНИИ")
    document.add_paragraph("Иванов И.И., Петров П.П.")
    document.add_paragraph("ТГТУ")
    document.add_paragraph("ivanov@example.com")
    document.add_paragraph("Научный руководитель: Сидоров С.С.")
    document.add_paragraph("Аннотация: В работе рассматривается учебное применение нейросетей.")
    document.add_paragraph("Ключевые слова: нейросети, образование, анализ данных")
    document.add_heading("Введение", level=1)
    document.add_paragraph("Цель работы — разработать демонстрационный программный модуль.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Показатель"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "Точность"
    table.cell(1, 1).text = "0.91"
    document.add_heading("Список литературы", level=1)
    document.add_paragraph("1. Иванов И.И. Нейросетевые методы. 2025.")
    document.save(buffer)
    return buffer.getvalue()


def build_docx_with_image() -> bytes:
    buffer = BytesIO()
    image_buffer = BytesIO()
    Image.new("RGB", (24, 24), "white").save(image_buffer, format="PNG")
    image_buffer.seek(0)

    document = Document()
    document.add_paragraph("ТЕСТОВЫЙ МАТЕРИАЛ С РИСУНКОМ")
    document.add_paragraph("Иванов И.И.")
    document.add_paragraph("ТГТУ")
    document.add_paragraph("ivanov@example.com")
    document.add_paragraph("Аннотация: Проверка сохранения изображения.")
    document.add_paragraph("Ключевые слова: изображение, тест")
    document.add_paragraph("Основной текст материала.")
    document.add_picture(image_buffer)
    document.add_paragraph("Рис. 1. Тестовый рисунок")
    document.save(buffer)
    return buffer.getvalue()


class ConferenceSystemTests(TestCase):
    def setUp(self):
        registry.clear()
        self.temp_dir = tempfile.TemporaryDirectory()
        self.override = override_settings(
            MEDIA_ROOT=Path(self.temp_dir.name),
            SUBMISSIONS_STORAGE_DIR=Path(self.temp_dir.name) / "submissions",
            CONTENT_VALIDATION_ENABLED=False,
        )
        self.override.enable()
        self.service = SubmissionService()
        self.service.ensure_defaults()

    def tearDown(self):
        registry.clear()
        self.override.disable()
        self.temp_dir.cleanup()

    def create_submission(self, content: bytes | None = None) -> dict:
        file = SimpleUploadedFile(
            "paper.docx",
            content if content is not None else build_valid_docx(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        return self.service.create_submission(
            {
                "author_contact": {
                    "full_name": "Тестовый Автор",
                    "email": "test@example.com",
                    "organization": "Новая тестовая организация",
                },
                "metadata": {
                    "title_ru": "Тестовый материал",
                    "section": "Искусственный интеллект в образовании",
                    "authors": [],
                },
            },
            uploaded_file=file,
        )

    def test_seed_organizations_loaded(self):
        self.assertTrue(Organization.objects.exists())

    def test_new_organization_is_saved(self):
        submission = self.create_submission()
        self.assertTrue(Submission.objects.filter(pk=submission["submission_id"]).exists())
        self.assertTrue(Organization.objects.filter(name="Новая тестовая организация").exists())

    def test_api_submissions_list(self):
        response = Client().get("/api/submissions/")
        self.assertEqual(response.status_code, 200)

    def test_docx_workflow_integrates_with_core(self):
        submission = self.create_submission()
        submission_id = submission["submission_id"]

        results = WorkflowEngine().run(submission_id)

        self.assertEqual(
            [row["stage_id"] for row in results],
            ["extract_metadata", "format_to_template", "content_validation", "export_pdf_and_package"],
        )
        self.assertNotIn("failed", [row["status"] for row in results])

        model = Submission.objects.get(pk=submission_id)
        self.assertEqual(model.status, "formatted")
        self.assertEqual(model.metadata["extracted_metadata"]["title"], "ПРИМЕНЕНИЕ НЕЙРОСЕТЕЙ В ОБРАЗОВАНИИ")
        self.assertTrue(model.metadata["body_text"])

        files = {
            row.file_type: row.path
            for row in SubmissionFile.objects.filter(submission_id=submission_id)
        }
        self.assertIn("extracted_metadata", files)
        self.assertIn("formatted_docx", files)
        self.assertIn("formatting_report", files)
        self.assertIn("formatted_pdf", files)
        self.assertIn("result_manifest", files)
        self.assertIn("result_package", files)
        self.assertTrue(Path(files["extracted_metadata"]).exists())
        self.assertTrue(Path(files["formatted_docx"]).exists())
        self.assertTrue(resolve_stored_file_path(files["formatting_report"]).exists())
        self.assertTrue(resolve_stored_file_path(files["formatted_pdf"]).exists())
        self.assertTrue(resolve_stored_file_path(files["result_manifest"]).exists())
        self.assertTrue(resolve_stored_file_path(files["result_package"]).exists())
        self.assertEqual(WorkflowStageResult.objects.filter(submission_id=submission_id).count(), 4)

    def test_second_stage_receives_file_created_by_first_stage(self):
        submission = self.create_submission()
        results = WorkflowEngine().run(submission["submission_id"])
        formatting = next(row for row in results if row["stage_id"] == "format_to_template")
        self.assertIn(formatting["status"], {"success", "warning"})
        self.assertTrue(formatting["result"].get("formatted_docx"))

    def test_images_are_saved_in_configured_shared_storage(self):
        submission = self.create_submission(content=build_docx_with_image())
        WorkflowEngine().run(submission["submission_id"])

        metadata_reference = self.service.get_latest_file_path(
            submission["submission_id"],
            "extracted_metadata",
        )
        metadata_path = resolve_stored_file_path(metadata_reference)
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        saved_images = metadata.get("objects", {}).get("saved_images", [])

        self.assertEqual(len(saved_images), 1)
        saved_path = Path(saved_images[0]["file_path"]).resolve()
        expected_root = (Path(self.temp_dir.name) / "submissions").resolve()
        self.assertTrue(saved_path.is_relative_to(expected_root))
        self.assertTrue(saved_path.exists())

    def test_invalid_docx_sets_error_and_skips_formatting(self):
        submission = self.create_submission(content=b"not a real docx")
        results = WorkflowEngine().run(submission["submission_id"])
        model = Submission.objects.get(pk=submission["submission_id"])

        self.assertEqual(results[0]["status"], "failed")
        self.assertEqual(results[1]["status"], "skipped")
        self.assertEqual(results[2]["status"], "skipped")
        self.assertEqual(results[3]["status"], "skipped")
        self.assertEqual(model.status, "error")


    @override_settings(
        CONTENT_VALIDATION_ENABLED=True,
        CONTENT_VALIDATION_API_URL="http://127.0.0.1:5100",
        CONTENT_VALIDATION_TIMEOUT=5,
    )
    def test_content_validation_stage_imports_results(self):
        submission = self.create_submission()
        submission_id = submission["submission_id"]
        submission_dir = Path(settings.SUBMISSIONS_STORAGE_DIR) / submission_id

        def fake_post(url, payload, timeout):
            self.assertEqual(payload, {"submissionId": submission_id})
            checks_dir = submission_dir / "checks"
            checks_dir.mkdir(parents=True, exist_ok=True)
            (checks_dir / "semantic_quality_check.json").write_text(
                json.dumps(
                    {
                        "check_id": "semantic_quality_check",
                        "title": "Смысловая проверка материала",
                        "status": "warning",
                        "risk_level": "medium",
                        "score": 0.7,
                        "summary": "Нужно уточнить методику.",
                        "warnings": [{"message": "Слабое описание метода."}],
                        "errors": [],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            (submission_dir / "check_result.json").write_text(
                json.dumps(
                    {
                        "submission_id": submission_id,
                        "overall_status": "needs_attention",
                        "overall_risk_level": "medium",
                        "author_message": "Уточните методику.",
                        "editor_message": "Проверьте методику.",
                        "checks": [
                            {
                                "check_id": "semantic_quality_check",
                                "title": "Смысловая проверка материала",
                                "status": "warning",
                                "risk_level": "medium",
                                "summary": "Нужно уточнить методику.",
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            return {"status": "success", "submission_id": submission_id}

        with patch(
            "submissions.integrations.content_validation_stage._post_json",
            side_effect=fake_post,
        ):
            results = WorkflowEngine().run(submission_id)

        validation = next(row for row in results if row["stage_id"] == "content_validation")
        self.assertEqual(validation["status"], "warning")
        self.assertEqual(Submission.objects.get(pk=submission_id).status, "needs_author_review")
        report_reference = self.service.get_latest_file_path(submission_id, "check_report")
        self.assertTrue(report_reference)
        self.assertEqual(resolve_stored_file_path(report_reference), submission_dir / "check_result.json")
        self.assertEqual(
            Submission.objects.get(pk=submission_id).checks.filter(
                check_id="semantic_quality_check"
            ).count(),
            1,
        )
